from datetime import datetime
import io
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

def generate_excel(df, batch_id, summary):
    output = io.BytesIO()

    # Required report
    required = pd.DataFrame({
        "Date (PDI Clearance)": [datetime_value for datetime_value in
            [pd.Timestamp.now().strftime("%Y-%m-%d")] * len(df)],
        "Chassis No": df["Chassis No"],
        "Bus No": df["Bus No"],
        "Vehicle Type": (
            df["Vehicle Type"]
            if "Vehicle Type" in df.columns
            else ""
        ),
        "PDI Status": df["PDI Status"],
        "QA Book Status": df["QA Book Status"],
    })

    summary_df = pd.DataFrame([
        ["Batch ID", batch_id],
        ["Observations", summary.get("observations", "")],
        ["Scanned QA Books", summary["scanned"]],
        ["PDI Cleared", summary["pdi_cleared"]],
        ["QA Books Received", summary["received"]],
    ], columns=["Metric", "Value"])

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        required.to_excel(writer, index=False, sheet_name="QA_Register")
        df.to_excel(writer, index=False, sheet_name="Detailed_Records")
        summary_df.to_excel(writer, index=False, sheet_name="Batch_Summary")

    output.seek(0)
    wb = load_workbook(output)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
        for col in ws.columns:
            max_len = 0
            for cell in col:
                max_len = max(max_len, len(str(cell.value or "")))
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 35)

    final = io.BytesIO()
    wb.save(final)
    final.seek(0)
    return final.getvalue()